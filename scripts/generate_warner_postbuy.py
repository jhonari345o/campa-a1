"""Genera el seed SQL, el catálogo TypeScript y las evidencias Warner 2026.

La hoja ``con_imagenes`` es la fuente visual. La hoja ``presentacion_FINAL``
se usa como control independiente de las filas y métricas. Los artes de
referencia cuya nota indica que la foto está pendiente no se publican como
evidencia.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import unicodedata
import uuid
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from PIL import Image, ImageOps


DEFAULT_VISUAL_BOOK = Path(r"C:\Users\jhonm\Downloads\Warner_medios_2026_con_imagenes.xlsx")
DEFAULT_CONTROL_BOOK = Path(r"C:\Users\jhonm\Downloads\Warner_medios_2026_presentacion_FINAL.xlsx")
DEFAULT_LOGOS = Path(r"C:\Users\jhonm\Downloads\Logos_radios_Warner_2026.zip")
UUID_NAMESPACE = uuid.UUID("5ae0597e-d5ae-467f-99e3-14ed251ba62c")
MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}
LOGOS = {
    "fm mundo": "/postbuy/warner/logos/fm-mundo.png",
    "blue radio": "/postbuy/warner/logos/blue-radio-1013.png",
    "fuego": "/postbuy/warner/logos/radio-fuego-1065.jpg",
}
ZIP_LOGOS = {
    "fm_mundo.png": "fm-mundo.png",
    "fm_mundo.svg": "fm-mundo.svg",
    "blue_radio_1013.png": "blue-radio-1013.png",
    "radio_fuego_1065.jpg": "radio-fuego-1065.jpg",
    "FUENTES.txt": "FUENTES.txt",
}


def slugify(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-") or "item"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sql_text(value: str | None) -> str:
    if value is None:
        return "null"
    return "'" + value.replace("'", "''") + "'"


def sql_number(value: int | float | None) -> str:
    if value is None:
        return "null"
    return str(value)


def integer(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    return None


def floating(value: Any) -> float | None:
    return float(value) if isinstance(value, (int, float)) else None


def stable_uuid(*parts: str) -> str:
    return str(uuid.uuid5(UUID_NAMESPACE, "|".join(parts)))


def media_prefix(placement: str) -> str:
    return placement.split("·", 1)[0].strip()


def channel_for(placement: str) -> str:
    prefix = media_prefix(placement).casefold()
    if prefix in LOGOS:
        return "radio"
    if prefix.startswith("arco"):
        return "btl"
    return "ooh"


def logo_for(placement: str) -> str | None:
    return LOGOS.get(media_prefix(placement).casefold())


def is_pending(note: str | None) -> bool:
    return bool(note and "pendiente" in note.casefold())


def extract_logo_files(zip_path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as archive:
        for source_name, output_name in ZIP_LOGOS.items():
            member = archive.getinfo(source_name)
            (destination / output_name).write_bytes(archive.read(member))


def save_evidence(image, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(io.BytesIO(image._data())) as source:
        converted = ImageOps.exif_transpose(source).convert("RGB")
        converted.thumbnail((1600, 1200), Image.Resampling.LANCZOS)
        converted.save(destination, "WEBP", quality=82, method=6)


def worksheet_max_row(sheet) -> int:
    """Obtiene el límite aun cuando openpyxl no lo precalcula en read-only."""
    if sheet.max_row is not None:
        return sheet.max_row
    _, _, _, max_row = range_boundaries(sheet.calculate_dimension(force=True))
    return max_row


def workbook_snapshot(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    result: dict[str, dict[str, Any]] = {}
    for sheet in workbook.worksheets:
        rows = []
        for row_index in range(8, worksheet_max_row(sheet) + 1):
            month = sheet.cell(row_index, 1).value
            if month == "Total":
                break
            campaign = sheet.cell(row_index, 2).value
            placement = sheet.cell(row_index, 3).value
            if not campaign or not placement:
                continue
            rows.append({
                "month": str(month).strip(),
                "campaign": str(campaign).strip(),
                "placement": str(placement).strip(),
                "quantity": integer(sheet.cell(row_index, 4).value),
                "rights": integer(sheet.cell(row_index, 5).value),
                "impacts": integer(sheet.cell(row_index, 6).value),
            })
        result[sheet.title] = {
            "month": str(sheet["B4"].value).strip(),
            "ooh_reach": floating(sheet["E4"].value),
            "radio_reach": floating(sheet["E5"].value) if sheet["D5"].value else None,
            "total_reach": floating(sheet["E6"].value) if sheet["D6"].value else None,
            "rows": rows,
        }
    return result


def verify_control_book(visual_path: Path, control_path: Path) -> None:
    visual = workbook_snapshot(visual_path)
    control = workbook_snapshot(control_path)
    if visual != control:
        visual_keys = set(visual)
        control_keys = set(control)
        raise ValueError(
            "La hoja de control no coincide con la fuente visual. "
            f"Solo visual: {sorted(visual_keys - control_keys)}; solo control: {sorted(control_keys - visual_keys)}"
        )


def read_campaigns(book_path: Path, evidence_root: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(book_path, data_only=True, read_only=False)
    campaigns: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        month_label = str(sheet["B4"].value).strip()
        month_number = MONTHS[month_label.casefold()]
        execution_month = f"2026-{month_number:02d}-01"
        campaign_name = str(sheet["A2"].value or sheet.title).strip()
        campaign_slug = slugify(campaign_name)
        campaign_id = stable_uuid("campaign", execution_month, campaign_name)
        images_by_row = {
            image.anchor._from.row + 1: image
            for image in sheet._images
            if getattr(image.anchor, "_from", None) is not None
        }
        placements: list[dict[str, Any]] = []
        for row_index in range(8, worksheet_max_row(sheet) + 1):
            if sheet.cell(row_index, 1).value == "Total":
                break
            placement_value = sheet.cell(row_index, 3).value
            if not placement_value:
                continue
            placement = str(placement_value).strip()
            note_value = sheet.cell(row_index, 8).value
            note = str(note_value).strip() if note_value else None
            channel = channel_for(placement)
            evidence_url: str | None = None
            if channel != "radio" and not is_pending(note) and row_index in images_by_row:
                relative = Path("evidence") / campaign_slug / f"{row_index:02d}-{slugify(placement)[:72]}.webp"
                save_evidence(images_by_row[row_index], evidence_root / relative)
                evidence_url = "/postbuy/warner/" + relative.as_posix()
            placement_id = stable_uuid("placement", campaign_id, str(row_index), placement)
            placements.append({
                "id": placement_id,
                "sourceRow": row_index,
                "name": placement,
                "channel": channel,
                "mediaName": media_prefix(placement),
                "quantity": integer(sheet.cell(row_index, 4).value),
                "mediaRights": integer(sheet.cell(row_index, 5).value),
                "monthlyImpacts": integer(sheet.cell(row_index, 6).value),
                "mediaLogoUrl": logo_for(placement),
                "evidenceImageUrl": evidence_url,
                "evidenceStatus": "verified" if evidence_url else "pending",
                "sourceNote": note,
            })
        evidence_count = sum(1 for item in placements if item["evidenceImageUrl"])
        status = "completed" if evidence_count == len(placements) else "partial" if evidence_count else "pending"
        campaigns.append({
            "id": campaign_id,
            "clientName": "Warner Bros. Discovery Ecuador",
            "executionMonth": execution_month,
            "monthLabel": month_label,
            "campaignName": campaign_name,
            "status": status,
            "oohReach": floating(sheet["E4"].value),
            "radioReach": floating(sheet["E5"].value) if sheet["D5"].value else None,
            "totalReach": floating(sheet["E6"].value) if sheet["D6"].value else None,
            "sourceFile": book_path.name,
            "sourceSheet": sheet.title,
            "placements": placements,
        })
    return campaigns


def type_script(campaigns: list[dict[str, Any]]) -> str:
    data = json.dumps(campaigns, ensure_ascii=False, indent=2)
    return f'''// Archivo generado por scripts/generate_warner_postbuy.py. No editar a mano.
export type PostbuyPlacement = {{
  id: string;
  sourceRow: number;
  name: string;
  channel: "ooh" | "radio" | "btl";
  mediaName: string;
  quantity: number | null;
  mediaRights: number | null;
  monthlyImpacts: number | null;
  mediaLogoUrl: string | null;
  evidenceImageUrl: string | null;
  evidenceStatus: "verified" | "pending";
  sourceNote: string | null;
}};

export type WarnerPostbuyCampaign = {{
  id: string;
  clientName: string;
  executionMonth: string;
  monthLabel: string;
  campaignName: string;
  status: "completed" | "partial" | "pending";
  oohReach: number | null;
  radioReach: number | null;
  totalReach: number | null;
  sourceFile: string;
  sourceSheet: string;
  placements: PostbuyPlacement[];
}};

export const WARNER_POSTBUY_CAMPAIGNS = {data} as const satisfies readonly WarnerPostbuyCampaign[];
'''


def seed_sql(campaigns: list[dict[str, Any]]) -> str:
    campaign_rows = []
    placement_rows = []
    for campaign in campaigns:
        campaign_rows.append(
            "(" + ", ".join([
                f"{sql_text(campaign['id'])}::uuid",
                "null",
                sql_text(campaign["clientName"]),
                f"{sql_text(campaign['executionMonth'])}::date",
                sql_text(campaign["campaignName"]),
                sql_text(campaign["status"]),
                sql_number(campaign["oohReach"]),
                sql_number(campaign["radioReach"]),
                sql_number(campaign["totalReach"]),
                sql_text(campaign["sourceFile"]),
                sql_text(campaign["sourceSheet"]),
            ]) + ")"
        )
        for position, item in enumerate(campaign["placements"], start=1):
            placement_rows.append(
                "(" + ", ".join([
                    f"{sql_text(item['id'])}::uuid",
                    f"{sql_text(campaign['id'])}::uuid",
                    sql_text(item["channel"]),
                    sql_text(item["mediaName"]),
                    sql_text(item["name"]),
                    sql_number(item["quantity"]),
                    sql_number(item["mediaRights"]),
                    sql_number(item["monthlyImpacts"]),
                    sql_text(item["evidenceImageUrl"]),
                    sql_text(item["mediaLogoUrl"]),
                    sql_text(item["evidenceStatus"]),
                    sql_text(item["sourceNote"]),
                    str(item["sourceRow"]),
                    str(position),
                ]) + ")"
            )
    return """-- Warner 2026 - campañas y ubicaciones de Reportería/Post-buys
-- Ejecutar después de supabase/migrations/0015_reporting_postbuys.sql.

begin;

insert into public.postbuy_campaigns
  (id, company_id, client_name, execution_month, campaign_name, status,
   ooh_reach, radio_reach, total_reach, source_file, source_sheet)
values
""" + ",\n".join(campaign_rows) + """
on conflict (id) do update set
  client_name = excluded.client_name,
  execution_month = excluded.execution_month,
  campaign_name = excluded.campaign_name,
  status = excluded.status,
  ooh_reach = excluded.ooh_reach,
  radio_reach = excluded.radio_reach,
  total_reach = excluded.total_reach,
  source_file = excluded.source_file,
  source_sheet = excluded.source_sheet,
  updated_at = now();

insert into public.postbuy_placements
  (id, campaign_id, channel, media_name, placement_name, element_quantity,
   media_rights, monthly_impacts, evidence_image_url, media_logo_url,
   evidence_status, source_note, source_row, display_order)
values
""" + ",\n".join(placement_rows) + """
on conflict (id) do update set
  campaign_id = excluded.campaign_id,
  channel = excluded.channel,
  media_name = excluded.media_name,
  placement_name = excluded.placement_name,
  element_quantity = excluded.element_quantity,
  media_rights = excluded.media_rights,
  monthly_impacts = excluded.monthly_impacts,
  evidence_image_url = excluded.evidence_image_url,
  media_logo_url = excluded.media_logo_url,
  evidence_status = excluded.evidence_status,
  source_note = excluded.source_note,
  source_row = excluded.source_row,
  display_order = excluded.display_order,
  updated_at = now();

commit;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-book", type=Path, default=DEFAULT_VISUAL_BOOK)
    parser.add_argument("--control-book", type=Path, default=DEFAULT_CONTROL_BOOK)
    parser.add_argument("--logos", type=Path, default=DEFAULT_LOGOS)
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    for path in (args.visual_book, args.control_book, args.logos):
        if not path.exists():
            raise SystemExit(f"No existe la fuente requerida: {path}")

    verify_control_book(args.visual_book, args.control_book)
    public_root = args.repo / "public" / "postbuy" / "warner"
    extract_logo_files(args.logos, public_root / "logos")
    campaigns = read_campaigns(args.visual_book, public_root)

    (args.repo / "lib" / "warner-postbuy-data.ts").write_text(type_script(campaigns), encoding="utf-8")
    import_root = args.repo / "supabase" / "imports" / "warner_2026"
    import_root.mkdir(parents=True, exist_ok=True)
    (import_root / "01_seed_warner_2026.sql").write_text(seed_sql(campaigns), encoding="utf-8")
    manifest = {
        "source": {
            "visual_file": args.visual_book.name,
            "visual_sha256": sha256(args.visual_book),
            "control_file": args.control_book.name,
            "control_sha256": sha256(args.control_book),
            "logos_file": args.logos.name,
            "logos_sha256": sha256(args.logos),
        },
        "campaigns": len(campaigns),
        "placements": sum(len(item["placements"]) for item in campaigns),
        "evidence_ready": sum(bool(row["evidenceImageUrl"]) for item in campaigns for row in item["placements"]),
        "evidence_pending": sum(not row["evidenceImageUrl"] for item in campaigns for row in item["placements"]),
    }
    (import_root / "IMPORT_MANIFEST.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
